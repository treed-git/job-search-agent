from __future__ import annotations

import io
import json
import logging
import re
from pathlib import Path

import fitz
from docx import Document as DocxDocument
from docx.shared import Pt, Inches
from fastapi import APIRouter, Request
from fastapi.responses import HTMLResponse, RedirectResponse, StreamingResponse
from fastapi.templating import Jinja2Templates

from app.config import settings, BASE_DIR
from app.database import get_job, save_generated_docs, update_job_status
from app.models import Resume, JobStatus
from app.ai.resume_tailor import tailor_resume
from app.ai.cover_letter import draft_cover_letter

logger = logging.getLogger(__name__)
router = APIRouter()
templates = Jinja2Templates(directory="app/templates")


def load_resume() -> Resume:
    path = Path(BASE_DIR) / settings.resume_path
    if not path.exists():
        logger.warning(f"Resume file not found at {path}")
        return Resume()
    with open(path) as f:
        return Resume(**json.load(f))


@router.post("/jobs/{job_id}/generate")
async def generate_documents(request: Request, job_id: int):
    """Generate a tailored resume and cover letter for an approved job."""
    job = await get_job(job_id)
    if not job:
        return RedirectResponse("/", status_code=302)

    resume = load_resume()

    tailored = await tailor_resume(resume, job.title, job.description)
    letter = await draft_cover_letter(resume, job.title, job.company, job.description)

    await save_generated_docs(job_id, tailored, letter)
    await update_job_status(job_id, JobStatus.APPROVED)

    return RedirectResponse(f"/jobs/{job_id}", status_code=302)


@router.get("/jobs/{job_id}/resume", response_class=HTMLResponse)
async def view_resume(request: Request, job_id: int):
    job = await get_job(job_id)
    if not job:
        return RedirectResponse("/", status_code=302)
    return templates.TemplateResponse(
        "document.html",
        {"request": request, "job": job, "doc_type": "resume", "content": job.tailored_resume},
    )


@router.get("/jobs/{job_id}/cover-letter", response_class=HTMLResponse)
async def view_cover_letter(request: Request, job_id: int):
    job = await get_job(job_id)
    if not job:
        return RedirectResponse("/", status_code=302)
    return templates.TemplateResponse(
        "document.html",
        {"request": request, "job": job, "doc_type": "cover_letter", "content": job.cover_letter},
    )


def _build_docx(text: str, title: str) -> io.BytesIO:
    """Convert plain/markdown text into a formatted Word document."""
    doc = DocxDocument()

    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)
    style.paragraph_format.space_after = Pt(4)

    section = doc.sections[0]
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)

    for line in text.splitlines():
        stripped = line.strip()

        # Markdown-style headings
        if stripped.startswith("### "):
            p = doc.add_paragraph()
            p.style = doc.styles["Heading 3"]
            p.text = stripped[4:]
        elif stripped.startswith("## "):
            p = doc.add_paragraph()
            p.style = doc.styles["Heading 2"]
            p.text = stripped[3:]
        elif stripped.startswith("# "):
            p = doc.add_paragraph()
            p.style = doc.styles["Heading 1"]
            p.text = stripped[2:]
        # ALL-CAPS section headers (common in generated resumes)
        elif stripped.isupper() and 3 <= len(stripped) <= 60:
            p = doc.add_paragraph()
            p.style = doc.styles["Heading 2"]
            p.text = stripped.title()
        # Bullet points
        elif stripped.startswith(("- ", "* ", "• ")):
            doc.add_paragraph(stripped[2:], style="List Bullet")
        # Horizontal rule / separator
        elif re.fullmatch(r"[-=_*]{3,}", stripped):
            continue
        # Blank line
        elif not stripped:
            doc.add_paragraph("")
        else:
            # Strip bold markdown markers for body text
            clean = re.sub(r"\*\*(.+?)\*\*", r"\1", stripped)
            doc.add_paragraph(clean)

    buf = io.BytesIO()
    doc.save(buf)
    buf.seek(0)
    return buf


def _safe_filename(text: str) -> str:
    """Create a filesystem-safe string from arbitrary text."""
    return re.sub(r"[^a-zA-Z0-9_-]", "_", text)[:60]


def _build_pdf(text: str, title: str) -> io.BytesIO:
    """Convert plain/markdown-like text into a basic PDF document."""
    pdf = fitz.open()
    page = pdf.new_page(width=612, height=792)  # US Letter

    margin = 54
    page_width = page.rect.width
    page_height = page.rect.height
    content_width = page_width - (2 * margin)
    y = margin

    title_height = 24
    title_rect = fitz.Rect(margin, y, page_width - margin, y + title_height)
    page.insert_textbox(title_rect, title, fontsize=16, fontname="helv", align=0)
    y += title_height + 12

    for raw_line in text.splitlines():
        line = raw_line.strip()
        if not line:
            y += 8
            continue

        if line.startswith("### "):
            line = line[4:]
            fontsize = 12
        elif line.startswith("## "):
            line = line[3:]
            fontsize = 13
        elif line.startswith("# "):
            line = line[2:]
            fontsize = 14
        elif line.startswith(("- ", "* ", "• ")):
            line = f"• {line[2:]}"
            fontsize = 11
        else:
            line = re.sub(r"\*\*(.+?)\*\*", r"\1", line)
            fontsize = 11

        if re.fullmatch(r"[-=_*]{3,}", line):
            continue

        rect = fitz.Rect(margin, y, margin + content_width, y + 200)
        required = page.insert_textbox(rect, line, fontsize=fontsize, fontname="helv", align=0)

        while required < 0:
            page = pdf.new_page(width=612, height=792)
            y = margin
            rect = fitz.Rect(margin, y, margin + content_width, y + 200)
            required = page.insert_textbox(rect, line, fontsize=fontsize, fontname="helv", align=0)

        consumed_height = 200 + required
        y += max(consumed_height + 4, fontsize + 6)

        if y > page_height - margin - 40:
            page = pdf.new_page(width=612, height=792)
            y = margin

    buf = io.BytesIO(pdf.tobytes())
    buf.seek(0)
    pdf.close()
    return buf


@router.get("/jobs/{job_id}/resume/download")
async def download_resume(job_id: int):
    """Download tailored resume as a Word document."""
    job = await get_job(job_id)
    if not job or not job.tailored_resume:
        return RedirectResponse(f"/jobs/{job_id}", status_code=302)

    buf = _build_docx(job.tailored_resume, f"Resume – {job.title}")
    filename = f"Resume_{_safe_filename(job.company)}_{_safe_filename(job.title)}.docx"
    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )


@router.get("/jobs/{job_id}/resume/download/pdf")
async def download_resume_pdf(job_id: int):
    """Download tailored resume as a PDF document."""
    job = await get_job(job_id)
    if not job or not job.tailored_resume:
        return RedirectResponse(f"/jobs/{job_id}", status_code=302)

    filename = f"Resume_{_safe_filename(job.company)}_{_safe_filename(job.title)}.pdf"
    buf = _build_pdf(job.tailored_resume, f"Resume - {job.title}")
    return StreamingResponse(
        buf,
        media_type="application/pdf",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )


@router.get("/jobs/{job_id}/cover-letter/download")
async def download_cover_letter(job_id: int):
    """Download cover letter as a Word document."""
    job = await get_job(job_id)
    if not job or not job.cover_letter:
        return RedirectResponse(f"/jobs/{job_id}", status_code=302)

    buf = _build_docx(job.cover_letter, f"Cover Letter – {job.title}")
    filename = f"Cover_Letter_{_safe_filename(job.company)}_{_safe_filename(job.title)}.docx"
    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )


@router.get("/jobs/{job_id}/cover-letter/download/pdf")
async def download_cover_letter_pdf(job_id: int):
    """Download cover letter as a PDF document."""
    job = await get_job(job_id)
    if not job or not job.cover_letter:
        return RedirectResponse(f"/jobs/{job_id}", status_code=302)

    filename = f"Cover_Letter_{_safe_filename(job.company)}_{_safe_filename(job.title)}.pdf"
    buf = _build_pdf(job.cover_letter, f"Cover Letter - {job.title}")
    return StreamingResponse(
        buf,
        media_type="application/pdf",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )
